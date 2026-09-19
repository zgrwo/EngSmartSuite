"""export_workbook 单元格/行截断回归测试（审查 2026-09-19 E13）。

修复前行为（双重静默截断，用户在表内看不到任何提示）：

1. 非数值单元格 `str(val)[:100]` 静默砍到 100 字符 —— 判定依据/结论建议等列
   常超 100 字符，审计工作簿因此丢失关键证据；
2. 数据行循环 `table.head(100)` 静默丢行（长表只留前 100 行，无标注）。

另：openpyxl 对超过 32767 字符的字符串在 save 时**静默**截到 32767（实测
无异常、无告警），因此超限文本必须显式截断并留下可见标记。

本文件用 canned orchestrate 结果直测导出层：不依赖真实引擎恰好产出长文本或长表。
"""

import os
import tempfile

import openpyxl
import pandas as pd
import pytest

from smartsuite.core.contracts import AnalysisResult
from smartsuite.services import audit as audit_module
from smartsuite.services.audit import export_workbook

# Excel 单元格字符上限（openpyxl 超限静默截断）
_EXCEL_CELL_MAX = 32767
# 每张表导出到工作簿的最大数据行数（预览上限，原 table.head(100) 语义）
_PREVIEW_ROWS = 100


@pytest.fixture()
def export_xlsx(monkeypatch):
    """导出 canned 结果并返回已加载的 workbook；临时文件在用例结束后清理。

    仅测试导出层：把 orchestrator.orchestrate 换成返回预置 AnalysisResult 的桩，
    这样表长与文本长度完全可控。
    """
    box: dict[str, AnalysisResult] = {}
    paths: list[str] = []

    def _fake_orchestrate(_req):
        return box["result"]

    monkeypatch.setattr(audit_module, "orchestrate", _fake_orchestrate)

    def _export(tables: dict[str, pd.DataFrame]):
        box["result"] = AnalysisResult(
            task="correlation", status="ok", summary="测试结论", tables=tables
        )
        fd, path = tempfile.mkstemp(suffix=".xlsx")
        os.close(fd)
        paths.append(path)
        export_workbook(
            pd.DataFrame({"y": [1.0, 2.0, 3.0]}),
            target_col="y",
            feature_cols=[],
            output_path=path,
            tasks=["correlation"],
        )
        return openpyxl.load_workbook(path)

    yield _export

    for p in paths:
        if os.path.exists(p):
            os.unlink(p)


def _texts(ws) -> list[str]:
    """工作表内所有字符串单元格（用于断言标注与文本完整性）。"""
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def test_long_text_not_truncated_at_100_chars(export_xlsx):
    """120 字符文本必须原样写入（修复前被 str(val)[:100] 砍掉 20 字符）。"""
    text = "温" * 120
    wb = export_xlsx({"main": pd.DataFrame({"判定依据": [text]})})
    try:
        ws = wb["correlation_summary"]
        assert text in _texts(ws), f"120 字符文本应完整保留，实际: {[t[:20] for t in _texts(ws)]}"
        assert not any(t.endswith("…") for t in _texts(ws)), "100 字符内不应出现任何截断标记"
    finally:
        wb.close()


def test_row_truncation_is_visible_in_sheet(export_xlsx):
    """150 行表：仍只导出前 100 行，但必须在表名处显式标注「预览」。"""
    wb = export_xlsx({"main": pd.DataFrame({"序号": [float(i) for i in range(150)]})})
    try:
        ws = wb["correlation_summary"]
        texts = _texts(ws)
        assert any("预览" in t for t in texts), f"行截断必须显式可见: {texts}"
        assert any("150" in t for t in texts), f"应标注原始行数: {texts}"
        data_rows = sum(
            1
            for row in ws.iter_rows()
            if isinstance(row[0].value, (int, float)) and not isinstance(row[0].value, bool)
        )
        assert data_rows == _PREVIEW_ROWS, f"应导出 {_PREVIEW_ROWS} 行预览，实际 {data_rows}"
    finally:
        wb.close()


def test_short_table_has_no_preview_annotation(export_xlsx):
    """未截断的表不得出现「预览」字样（防标注滥发）。"""
    wb = export_xlsx({"main": pd.DataFrame({"序号": [float(i) for i in range(5)]})})
    try:
        texts = _texts(wb["correlation_summary"])
        assert not any("预览" in t for t in texts), f"短表不应标注预览: {texts}"
    finally:
        wb.close()


def test_cell_over_excel_limit_truncated_with_marker(export_xlsx):
    """超 Excel 单元格上限（32767）的长文本：按上限截断并显式标注。"""
    text = "长" * 40000
    wb = export_xlsx({"main": pd.DataFrame({"备注": [text]})})
    try:
        long_cells = [t for t in _texts(wb["correlation_summary"]) if len(t) > 1000]
        assert len(long_cells) == 1, f"应有且仅有一个长文本单元格: {[len(t) for t in long_cells]}"
        val = long_cells[0]
        assert len(val) <= _EXCEL_CELL_MAX, f"不得超过 Excel 单元格上限: {len(val)}"
        assert val.endswith("（已截断）"), f"截断必须显式标注: {val[-12:]!r}"
        assert val.startswith("长"), "应保留原文前缀"
    finally:
        wb.close()
